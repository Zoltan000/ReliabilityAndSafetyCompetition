"""Build the 2 Req. 1 slides (no title slide; all names on each slide) -> outputs/Req1_slides.pptx.

Usage: .venv/Scripts/python scripts/08_make_slides.py --names "Name One, Name Two, Name Three"
"""
import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

FIG = Path("outputs/figures")
INK, INK2 = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E)

SLIDES = [
    {
        "headline": "Bearings, belts and motor-reducers wear out, and load (not weather) sets how fast",
        "left": ("beta.png", "Failure behaviour: 3 wear-out, 2 random, software infant mortality"),
        "right": ("load.png", "Load is a stress multiplier, not just more hours"),
        "bullets": [
            "Bearings cause 94% of failures (≈41 per conveyor-year); every failure costs exactly 36 h of downtime",
            "Every component renews perfectly: first-life and replacement lives are identical. The apparent "
            "'shorter replacement belts' is Simpson's paradox (replacements happen mostly on Heavy conveyors)",
            "Weather does not drive failures: seasonal temperature steps and daily temp./humidity/voltage swings change "
            "failure rates by ≤ 5% (bearings: 0%)",
        ],
    },
    {
        "headline": "Downtime is a failure-count problem, and it is forecastable from the last observed day",
        "left": ("cm.png", "Motor-reducer and belt failures announce themselves; bearing failures do not"),
        "right": ("accuracy.png", "Held-out conveyors (grouped CV): 3-year downtime error"),
        "bullets": [
            "Downtime = 36 h × failures + 24 h × planned stops (3/yr, 15 Mar / 15 Jul / 15 Dec): forecasting it means "
            "forecasting the failure count",
            "Young conveyors ramp up: all bearings start new (β≈3), so failures rise through year 1 before settling. "
            "Tracking every bearing's age captures this",
            "Bad actors stay bad: a conveyor's failure rate vs. its load-class peers in years 1-10 predicts years 11-20 "
            "(r = 0.93), so each conveyor's own history is a key forecast input",
            "Contactor wear-out is emerging (12 failures, all on a restart after repair; β≈3.7, wide CI): watch it",
        ],
    },
]


def add_text(slide, x, y, w, h, text, size, color=INK, bold=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    items = text if isinstance(text, list) else [text]
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " + t) if isinstance(text, list) else t
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
        p.space_after = Pt(4)
    return tb


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--names", default="TEAM MEMBER NAMES")
    a = ap.parse_args()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W = prs.slide_width
    for s in SLIDES:
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        add_text(sl, Inches(0.5), Inches(0.3), W - Inches(1), Inches(0.9), s["headline"], 26, bold=True)
        for k, (img, cap) in enumerate([s["left"], s["right"]]):
            x = Inches(0.5) + k * Inches(6.3)
            add_text(sl, x, Inches(1.25), Inches(6.1), Inches(0.4), cap, 13, INK2, bold=True)
            if (FIG / img).exists():
                sl.shapes.add_picture(str(FIG / img), x, Inches(1.7), width=Inches(6.1))
        add_text(sl, Inches(0.5), Inches(5.3), W - Inches(1), Inches(1.6), s["bullets"], 14)
        add_text(sl, Inches(0.5), Inches(7.0), W - Inches(1), Inches(0.4),
                 f"{a.names}  ·  Industrial Conveyor Reliability Data Challenge  ·  fleet: 284 conveyors × 20 years",
                 10, INK2)
    out = Path("outputs/Req1_slides.pptx")
    prs.save(out)
    print(f"wrote {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
